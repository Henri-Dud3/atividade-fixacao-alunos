"""Gera o documento de entrega com o link real do repositório."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(14)
doc.add_paragraph('Sistema de cadastro de alunos e notas', 'Title')
doc.add_paragraph('Atividade de fixação', 'Subtitle')
doc.add_paragraph('O sistema Nota Clara permite ao professor cadastrar alunos e suas notas finais. Cada aluno possui um acesso individual para consultar somente a própria nota.')
doc.add_heading('Link do projeto', 1)
p = doc.add_paragraph()
url = 'https://nota-clara-production.up.railway.app'
link = OxmlElement('w:hyperlink')
link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
run = OxmlElement('w:r')
props = OxmlElement('w:rPr')
color = OxmlElement('w:color'); color.set(qn('w:val'), '176B53'); props.append(color)
run.append(props)
text = OxmlElement('w:t'); text.text = url; run.append(text); link.append(run); p._p.append(link)
doc.add_paragraph('Acesse o sistema pelo link acima e utilize uma das contas abaixo. O banco de dados fica armazenado em um volume persistente no Railway.')
doc.add_paragraph('Código-fonte e instruções: https://github.com/Henri-Dud3/atividade-fixacao-alunos')
doc.add_heading('Usuários e senhas para avaliação', 1)
table = doc.add_table(rows=1, cols=4)
table.style = 'Light Shading Accent 1'
for cell, value in zip(table.rows[0].cells, ['Perfil', 'Nome', 'Usuário', 'Senha']): cell.text = value
for row in [('Professor', 'Professor', 'professor', 'Professor@123'), ('Aluno', 'Ana Silva', 'ana', 'Ana@12345'), ('Aluno', 'Bruno Santos', 'bruno', 'Bruno@12345')]:
    for cell, value in zip(table.add_row().cells, row): cell.text = value
for cell in table.rows[0].cells:
    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E8EDE9'); cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs: run.font.color.rgb = RGBColor(0, 0, 0); run.bold = True
doc.add_heading('Como executar', 1)
doc.add_paragraph('Com PHP 8.1 ou superior e a extensão PDO SQLite habilitada, abra o terminal na pasta do projeto e execute:')
p = doc.add_paragraph('php -S localhost:8000 -t public'); p.runs[0].font.name = 'Consolas'; p.runs[0].font.size = Pt(10)
doc.add_paragraph('Acesse http://localhost:8000 no navegador. O banco e as contas acima são criados automaticamente no primeiro acesso. No Windows preparado para esta atividade, também é possível executar o arquivo iniciar.ps1.')
doc.add_heading('Recursos e controle de acesso', 1)
doc.add_paragraph('O professor pode cadastrar, buscar, editar e excluir alunos. O aluno pode apenas consultar sua nota. Há uma conta ativa por sessão do navegador: para trocar de usuário, é necessário sair. As notas aceitam valores de 0 a 10.')
doc.add_paragraph('Tecnologias: HTML, CSS, JavaScript, PHP e SQL com SQLite. O projeto utiliza sessões, senhas protegidas por hash, consultas parametrizadas e validação no servidor.')
doc.core_properties.title = 'Sistema de cadastro de alunos e notas'
doc.core_properties.subject = 'Atividade de fixação'
doc.save(Path(__file__).with_name('Entrega_Railway.docx'))
